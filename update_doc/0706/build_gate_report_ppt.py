import sys
sys.path.insert(0, "/tmp/pptx_deps")

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path

OUT = Path(__file__).with_name("gate_layer_selection_report.pptx")
IMG = Path(__file__).with_name("gate_layer_selection_slide.png")

NAVY = RGBColor(15, 23, 42)
SLATE = RGBColor(71, 85, 105)
MUTED = RGBColor(148, 163, 184)
LIGHT = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(8, 145, 178)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(217, 119, 6)
RED = RGBColor(220, 38, 38)
PURPLE = RGBColor(124, 58, 237)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line or fill
    return shp

def text(slide, s, x, y, w, h, size=20, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         font="Noto Sans CJK SC", valign=MSO_ANCHOR.TOP, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.text = s; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    return box

def rich(slide, lines, x, y, w, h, size=18, color=NAVY, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.06)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = "Noto Sans CJK SC"; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(10); p.level = 0
        if bullet: p.text = "•  " + line
    return box

def title(slide, heading, kicker=None, page=None):
    if kicker: text(slide, kicker.upper(), .65, .25, 7, .3, 10, BLUE, True)
    text(slide, heading, .65, .58, 11.8, .62, 27, NAVY, True)
    rect(slide, .65, 1.25, .62, .05, BLUE)
    if page is not None: text(slide, f"{page:02d}", 12.25, .3, .45, .3, 10, MUTED, True, PP_ALIGN.RIGHT)

def footer(slide, label="Gate layer selection · 0705 update"):
    text(slide, label, .65, 7.15, 7, .2, 9, MUTED)

def metric_card(slide, x, y, w, label, old, new, delta, good=True):
    rect(slide, x, y, w, 1.35, WHITE, True, RGBColor(226,232,240))
    text(slide, label, x+.18, y+.14, w-.36, .28, 12, SLATE, True)
    text(slide, f"{old}  →  {new}", x+.18, y+.48, w-.36, .42, 21, NAVY, True)
    text(slide, delta, x+.18, y+.98, w-.36, .22, 11, GREEN if good else RED, True)

def add_chart(slide, categories, series, x, y, w, h, colors, ymin=None, ymax=None, legend=True, labels=True):
    data = CategoryChartData(); data.categories = categories
    for name, vals in series: data.add_series(name, vals)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_title = False; chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11); chart.legend.font.name = "Noto Sans CJK SC"
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(226,232,240)
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.size = Pt(11)
    if ymin is not None: chart.value_axis.minimum_scale = ymin
    if ymax is not None: chart.value_axis.maximum_scale = ymax
    for i, ser in enumerate(chart.series):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = colors[i]
        ser.format.line.color.rgb = colors[i]
        if labels:
            ser.has_data_labels = True; ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            ser.data_labels.font.size = Pt(9); ser.data_labels.number_format = '0.0000'
    return chart

# 1 cover
s = prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,13.333,7.5,LIGHT)
rect(s,0,0,4.25,7.5,NAVY)
text(s,"GATE\nLAYER\nSELECTION",.65,.75,3.0,2.25,34,WHITE,True)
text(s,"实验汇报",.68,3.35,2.8,.4,16,RGBColor(125,211,252),True)
text(s,"从统一末层到按 transition 选层",4.9,1.3,7.6,.85,30,NAVY,True)
text(s,"全量验证结论、pilot 探索过程与上线决策",4.92,2.25,6.8,.45,18,SLATE)
rect(s,4.92,3.2,6.85,1.35,WHITE,True,RGBColor(226,232,240))
text(s,"最终建议",5.2,3.45,1.4,.3,12,BLUE,True)
text(s,"j=0,1 → layer 15   |   j=2,3 → layer 23   |   LR",5.2,3.82,6.2,.42,20,NAVY,True)
text(s,"基于 83,927 train + 8,821 dev",4.92,5.65,5,.3,12,MUTED)
text(s,"0705 UPDATE",10.65,6.65,1.7,.25,10,BLUE,True,PP_ALIGN.RIGHT)

# 2 question
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"为什么要重新选择 gate 层？","背景",2); footer(s)
text(s,"现状",.7,1.62,1,.3,13,SLATE,True)
rect(s,.7,2.0,3.55,3.75,NAVY,True)
text(s,"Layer 31",1.0,2.42,2.9,.55,32,WHITE,True,PP_ALIGN.CENTER)
text(s,"PCA(64) + Logistic Regression",1.0,3.18,2.9,.55,15,RGBColor(203,213,225),False,PP_ALIGN.CENTER)
text(s,"4 个 pooled transition\n共享同一层",1.0,4.18,2.9,.85,18,WHITE,True,PP_ALIGN.CENTER)
text(s,"两个核心问题",4.85,1.62,2.5,.3,13,SLATE,True)
for i,(n,h,b,c) in enumerate([
    ("01","最后一层真的是最优层吗？","末层更贴近 next-token 目标，未必保留最强的连贯性信号。",BLUE),
    ("02","能否补强深层 transition？","生产 dev 中 K=4 的 F1 明显低于 K=2，是已知短板。",PURPLE)]):
    y=2.0+i*1.9; rect(s,4.85,y,7.7,1.55,WHITE,True,RGBColor(226,232,240))
    text(s,n,5.1,y+.28,.55,.45,18,c,True); text(s,h,5.85,y+.22,6.2,.35,19,NAVY,True)
    text(s,b,5.85,y+.75,6.15,.5,14,SLATE)
rect(s,4.85,5.95,7.7,.55,RGBColor(219,234,254),True)
text(s,"目标：提升判别力，同时不以更高 FPR 为代价。",5.15,6.09,7.0,.25,14,BLUE,True)

# 3 executive summary
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"全量验证：混层 + LR 全面优于生产现状","最终结论",3); footer(s)
metric_card(s,.7,1.65,2.32,"AUC","0.9074","0.9221","+0.0147")
metric_card(s,3.18,1.65,2.32,"TPR","0.8235","0.8456","+0.0221")
metric_card(s,5.66,1.65,2.32,"FPR","0.1756","0.1693","−0.0063 · 更低")
metric_card(s,8.14,1.65,2.32,"Precision","0.7084","0.7212","+0.0128")
metric_card(s,10.62,1.65,2.02,"F1","0.7616","0.7785","+0.0169")
rect(s,.7,3.45,12.0,2.25,NAVY,True)
text(s,"推荐配置",1.0,3.82,1.5,.3,12,RGBColor(125,211,252),True)
text(s,"Layer 15",1.0,4.35,2.0,.45,24,WHITE,True); text(s,"j=0 / j=1",1.0,4.86,2,.3,14,RGBColor(203,213,225))
text(s,"＋",3.15,4.42,.4,.4,24,RGBColor(125,211,252),True,PP_ALIGN.CENTER)
text(s,"Layer 23",3.75,4.35,2.0,.45,24,WHITE,True); text(s,"j=2 / j=3",3.75,4.86,2,.3,14,RGBColor(203,213,225))
text(s,"＋",5.95,4.42,.4,.4,24,RGBColor(125,211,252),True,PP_ALIGN.CENTER)
text(s,"Logistic Regression",6.55,4.35,3.0,.45,24,WHITE,True); text(s,"阈值校准稳定",6.55,4.86,2.4,.3,14,RGBColor(203,213,225))
rect(s,9.95,3.87,2.25,1.35,RGBColor(30,64,175),True)
text(s,"83,927",10.1,4.07,1.95,.42,25,WHITE,True,PP_ALIGN.CENTER); text(s,"全量 train",10.1,4.57,1.95,.25,12,RGBColor(191,219,254),False,PP_ALIGN.CENTER)
text(s,"结论不是用 FPR 换 TPR：TPR、Precision、F1 均提升，FPR 反而下降。",.8,6.18,11.8,.35,15,GREEN,True,PP_ALIGN.CENTER)

# 4 per j
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"提升主要来自 j=0 与 j=3","分 transition",4); footer(s)
add_chart(s,["j=0\nQ→E1","j=1\nE1→E2","j=2\nE2→E3","j=3\nE3→E4"],
          [("Layer 31 + LR",[.7513,.7699,.7610,.7843]),("混层 + LR",[.7876,.7760,.7535,.8056])],
          .7,1.55,8.2,4.85,[MUTED,BLUE],.70,.83)
rect(s,9.25,1.65,3.35,1.15,RGBColor(220,252,231),True)
text(s,"j=0",9.52,1.86,.6,.3,14,GREEN,True); text(s,"+0.0363",10.25,1.8,1.95,.45,23,GREEN,True)
rect(s,9.25,3.05,3.35,1.15,RGBColor(219,234,254),True)
text(s,"j=3",9.52,3.26,.6,.3,14,BLUE,True); text(s,"+0.0213",10.25,3.2,1.95,.45,23,BLUE,True)
rect(s,9.25,4.45,3.35,1.35,RGBColor(254,242,242),True)
text(s,"唯一回落",9.52,4.65,1.0,.3,12,RED,True); text(s,"j=2  −0.0075",9.52,5.08,2.5,.4,18,NAVY,True)
text(s,"整体收益覆盖了局部小幅回落；j=3 的改善直接对应深层 transition 短板。",9.28,6.05,3.2,.65,13,SLATE)

# 5 K
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"K=2 / 3 / 4 均提升，深层短板得到缓解","分 hop 数",5); footer(s)
add_chart(s,["K=2","K=3","K=4"],[("Layer 31 + LR",[.8069,.7467,.7138]),("混层 + LR",[.8265,.7638,.7268])],
          .7,1.55,8.3,4.85,[MUTED,CYAN],.68,.85)
text(s,"Δ F1",9.55,1.75,2.4,.4,16,SLATE,True,PP_ALIGN.CENTER)
for i,(k,d) in enumerate([("K=2","+0.0196"),("K=3","+0.0171"),("K=4","+0.0130")]):
    y=2.35+i*1.08; rect(s,9.45,y,2.65,.78,WHITE,True,RGBColor(207,250,254))
    text(s,k,9.7,y+.22,.7,.25,13,SLATE,True); text(s,d,10.45,y+.16,1.4,.35,18,CYAN,True,PP_ALIGN.RIGHT)
rect(s,9.28,5.78,3.0,.62,RGBColor(236,254,255),True)
text(s,"K=4：0.7138 → 0.7268",9.45,5.96,2.7,.25,13,CYAN,True,PP_ALIGN.CENTER)

# 6 why LR
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"为什么最终仍选择 LR，而不是 MLP？","模型选择",6); footer(s)
rect(s,.7,1.6,5.75,4.75,WHITE,True,RGBColor(226,232,240)); rect(s,6.85,1.6,5.75,4.75,WHITE,True,RGBColor(226,232,240))
text(s,"MLP · pilot",1.05,1.98,2.5,.4,21,PURPLE,True)
text(s,"AUC 更高",1.05,2.62,2.0,.35,17,NAVY,True); text(s,"0.9266",3.8,2.48,2.0,.5,28,PURPLE,True,PP_ALIGN.RIGHT)
text(s,"但阈值迁移失稳",1.05,3.38,3.0,.35,17,NAVY,True)
text(s,"目标 FPR 0.15",1.05,4.05,2.2,.3,14,SLATE); text(s,"dev 实际 0.2348",3.55,3.97,2.3,.42,22,RED,True,PP_ALIGN.RIGHT)
rect(s,1.05,4.76,4.9,.85,RGBColor(254,226,226),True)
text(s,"排序质量提升真实，但触发阈值尚不可控",1.28,5.0,4.45,.35,14,RED,True,PP_ALIGN.CENTER)
text(s,"LR · 全量",7.2,1.98,2.5,.4,21,BLUE,True)
text(s,"AUC 稳健提升",7.2,2.62,2.4,.35,17,NAVY,True); text(s,"+0.0147",10.1,2.48,1.9,.5,28,BLUE,True,PP_ALIGN.RIGHT)
text(s,"阈值校准可靠",7.2,3.38,2.5,.35,17,NAVY,True)
text(s,"生产 FPR 0.1756",7.2,4.05,2.5,.3,14,SLATE); text(s,"新方案 0.1693",9.85,3.97,2.15,.42,22,GREEN,True,PP_ALIGN.RIGHT)
rect(s,7.2,4.76,4.9,.85,RGBColor(220,252,231),True)
text(s,"性能提升且风险更低，适合进入下游验证",7.43,5.0,4.45,.35,14,GREEN,True,PP_ALIGN.CENTER)
text(s,"决策原则：当前先吃到“选层”的确定性收益；MLP 待重新设计校准方案后再评估。",.9,6.62,11.5,.35,15,NAVY,True,PP_ALIGN.CENTER)

# 7 exploration
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"pilot 如何找到混层方案？","探索路径",7); footer(s)
steps=[("1","扫描 7 / 15 / 23 / 31","中间层 15、23 整体优于末层 31",BLUE),
       ("2","按 transition 分解","15 擅长 j=0/1；23 擅长 j=2/3",CYAN),
       ("3","concat 对照","15+15 ≈ 15，排除“维度越多越好”",AMBER),
       ("4","加入 Layer 31","15+23+31 反而降至 0.9010",RED)]
for i,(n,h,b,c) in enumerate(steps):
    x=.75+i*3.08; rect(s,x,1.75,2.72,3.72,WHITE,True,RGBColor(226,232,240))
    rect(s,x+.2,2.0,.58,.58,c,True); text(s,n,x+.2,2.1,.58,.28,16,WHITE,True,PP_ALIGN.CENTER)
    text(s,h,x+.2,2.82,2.3,.72,18,NAVY,True)
    text(s,b,x+.2,3.85,2.3,1.15,14,SLATE)
text(s,"推断",.82,5.92,.75,.28,12,PURPLE,True)
text(s,"不同深度保留的信号具有互补性；最后一层可能更受 next-token 目标塑形。",1.65,5.85,10.8,.45,16,NAVY,True)
text(s,"注：这是基于实验结果的解释性推断，不是因果证明。",1.65,6.37,10.8,.3,11,MUTED)

# 8 engineering
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"方案已具备 artifact 驱动的工程落地能力","实现状态",8); footer(s)
items=[("数据", "hidden_states/pilot_multilayer/", "同一份多层 NPZ；train 83,927 / dev 8,821", BLUE),
       ("训练", "gate/artifacts_pooled_v2/", "每个 joblib 自带 layer 元数据", CYAN),
       ("训练脚本", "fit_lr_gate_pooled.py", "新增 --per-j-layers 与 split 参数", PURPLE),
       ("检索脚本", "run_retrieval_exp_wavefront.py", "一次 forward 读取所需层；按 artifact 选层", GREEN)]
for i,(a,b,c,d) in enumerate(items):
    y=1.55+i*1.17; rect(s,.75,y,11.8,.9,WHITE,True,RGBColor(226,232,240))
    rect(s,.75,y,.17,.9,d); text(s,a,1.12,y+.2,1.0,.3,14,d,True)
    text(s,b,2.35,y+.17,4.1,.35,15,NAVY,True,font="Noto Sans Mono CJK SC")
    text(s,c,6.65,y+.18,5.55,.38,14,SLATE)
rect(s,.75,6.35,11.8,.48,RGBColor(239,246,255),True)
text(s,"检索代码没有写死层号：配置由 artifact 自身元数据驱动。",1.0,6.48,11.3,.22,13,BLUE,True,PP_ALIGN.CENTER)

# 9 decision
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"上线前只剩一块关键拼图：端到端检索验证","下一步",9); footer(s)
rect(s,.75,1.55,7.65,4.95,NAVY,True)
text(s,"建议决策流程",1.1,1.95,2.5,.35,17,RGBColor(125,211,252),True)
for i,(n,h,b) in enumerate([("01","保持生产 artifact 不变","先不直接替换 gate/artifacts_pooled/"),
                            ("02","运行 A/B 端到端对比","v1 vs v2：recall / EM / F1"),
                            ("03","指标确认后转正","门控指标已通过，只需验证下游收益")]):
    y=2.55+i*1.08; rect(s,1.1,y,.52,.52,RGBColor(30,64,175),True)
    text(s,n,1.1,y+.1,.52,.22,12,WHITE,True,PP_ALIGN.CENTER)
    text(s,h,1.85,y-.02,2.7,.35,17,WHITE,True)
    text(s,b,4.75,y,3.05,.48,12,RGBColor(203,213,225))
rect(s,8.8,1.55,3.8,2.15,RGBColor(220,252,231),True)
text(s,"GO 条件",9.15,1.9,1.5,.3,14,GREEN,True)
text(s,"下游 recall / EM / F1\n至少不退化，并出现可复现收益",9.15,2.46,3.0,.8,18,NAVY,True)
rect(s,8.8,4.05,3.8,2.45,RGBColor(255,247,237),True)
text(s,"暂缓项",9.15,4.4,1.5,.3,14,AMBER,True)
rich(s,["MLP：重做阈值校准后再测","27 / 29 层：可选甜蜜点探索"],9.15,4.92,3.0,1.05,14,SLATE,True)

# 10 appendix image
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"原始实验总结图","附录",10); footer(s,"Source: update_doc/0706/gate_layer_selection_slide.png")
if IMG.exists():
    s.shapes.add_picture(str(IMG), Inches(.8), Inches(1.47), width=Inches(11.75), height=Inches(5.37))

prs.core_properties.title = "Gate 层选择 / 分类器 pilot 实验汇报"
prs.core_properties.subject = "基于 0705 update 报告生成"
prs.core_properties.author = "OpenAI Codex"
prs.core_properties.keywords = "gate, layer selection, logistic regression, retrieval"
prs.save(OUT)
print(OUT)
